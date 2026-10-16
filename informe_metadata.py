"""
Toda la ficha de cada libro en un Excel: autor, editorial, medidas y peso.

Pedido por operaciones el 16/10/2026. La ficha esta repartida en tres
origenes dentro de odoo_books_mirror y nadie la habia sacado junta:

    cdl_*       el scrapeo de Casa del Libro. Es el unico que trae PESO,
                ALTO y ANCHO, y cubre 355.412 libros.
    gbooks_*    Google Books: editorial, idioma y paginas. 282.427.
    azeta_*     precio y tipo de IVA del catalogo de AZETA.

Las medidas llegan como texto con su unidad -'502.0 gr', '19.1 cm'- asi
que aqui se pasan a numero: en Excel un peso que es texto no se puede
sumar ni filtrar por rango, que es justo para lo que sirve tenerlo.

Por defecto solo salen los libros que tienen ALGO -scrapeo o Google-,
que son unos 400.000 de los 1.315.812 del espejo. Con todos=true salen
los 1,3 millones, repartidos en varias hojas porque Excel no admite mas
de 1.048.576 filas por hoja.
"""
import os
import re
import time
from datetime import datetime

import db
from informe_stock import get_letra, _podar

CARPETA = os.environ.get("INFORME_STOCK_DIR", "/tmp")
TOPE_HOJA = 1_000_000

_job: dict | None = None


def get_status() -> dict:
    return dict(_job) if _job else {"status": "idle"}


COLUMNAS = [
    ("ISBN", 15),
    ("TITULO", 46),
    ("AUTOR", 28),
    ("EDITORIAL", 24),
    ("COLECCION", 22),
    ("TRADUCTOR", 20),
    ("ILUSTRADOR", 20),
    ("IDIOMA", 12),
    ("ENCUADERNACION", 16),
    ("PAGINAS", 9),
    ("PUBLICADO", 12),
    ("PESO (g)", 10),
    ("ALTO (cm)", 10),
    ("ANCHO (cm)", 11),
    ("PVP", 10),
    ("PVP SCRAPEADO", 13),
    ("CATEGORIA", 26),
    ("CATEGORIAS DEDUCIDAS", 30),
    ("DESCRIPCION (caracteres)", 13),
    ("PORTADA", 46),
    ("PROVEEDORES", 34),
    ("DE DONDE SALE LA FICHA", 18),
    ("ULTIMA ACTUALIZACION", 18),
]

AYUDA = {
    "ISBN": "El codigo del libro.",
    "TITULO": "Como figura en Odoo. Si aparece el propio ISBN es que la "
              "ficha nunca se enriquecio.",
    "AUTOR": "Del scrapeo de Casa del Libro.",
    "EDITORIAL": "Del scrapeo; si falta, la de Google Books.",
    "IDIOMA": "Del scrapeo; si falta, el de Google Books.",
    "PAGINAS": "Del scrapeo; si falta, las de Google Books.",
    "PESO (g)": "En gramos, ya como numero. Viene del scrapeo y es el unico "
                "sitio donde lo tenemos: Odoo lo guarda en kilos y casi "
                "siempre a cero.",
    "ALTO (cm)": "En centimetros, como numero.",
    "ANCHO (cm)": "En centimetros, como numero.",
    "PVP": "El precio de venta que tiene hoy en Odoo.",
    "PVP SCRAPEADO": "El PVP que se leyo de Casa del Libro el dia del alta. "
                     "Puede estar viejo: no se refresca.",
    "CATEGORIAS DEDUCIDAS": "Las que infiere el sistema cuando Odoo no trae "
                            "categoria.",
    "DESCRIPCION (caracteres)": "Cuanta sinopsis hay. Por debajo de 120 la "
                                "ficha se considera coja y no se publica en "
                                "la tienda.",
    "PROVEEDORES": "Quienes constan en Odoo como capaces de servirlo. No "
                   "significa que lo tengan ahora.",
    "DE DONDE SALE LA FICHA": "«scrapeo» Casa del Libro · «google» Google "
                              "Books · «ambos» · «ninguno».",
}

# '502.0 gr' -> 502.0 ; '19,1 cm' -> 19.1 ; '' -> None
_NUM = re.compile(r"-?\d+(?:[.,]\d+)?")


def _numero(txt):
    if txt is None:
        return None
    m = _NUM.search(str(txt))
    if not m:
        return None
    try:
        return round(float(m.group(0).replace(",", ".")), 2)
    except ValueError:
        return None


def _entero(txt):
    v = _numero(txt)
    return int(v) if v is not None else None


SQL = """
    SELECT m.barcode,
           COALESCE(NULLIF(m.name, ''), ''),
           m.cdl_author, m.cdl_editorial, m.cdl_collection,
           m.cdl_translator, m.cdl_illustrator,
           COALESCE(NULLIF(m.cdl_language, ''), m.gbooks_language),
           m.cdl_binding,
           COALESCE(NULLIF(m.cdl_pages, ''), m.gbooks_pages::text),
           m.cdl_release_date,
           m.cdl_weight, m.cdl_height, m.cdl_width,
           m.list_price, m.pvp_base,
           COALESCE(NULLIF(m.categ_name, ''), ''),
           m.inferred_categories,
           LENGTH(COALESCE(m.description, '')),
           COALESCE(NULLIF(m.cdl_image_url, ''), m.gbooks_thumbnail),
           m.supplier_names,
           COALESCE(NULLIF(m.gbooks_publisher, ''), ''),
           (m.cdl_fetched_at IS NOT NULL),
           (m.gbooks_fetched_at IS NOT NULL),
           GREATEST(m.cdl_fetched_at, m.gbooks_fetched_at, m.synced_at)
    FROM odoo_books_mirror m
    WHERE m.barcode IS NOT NULL AND m.barcode <> ''
"""
SOLO_CON_FICHA = ("  AND (m.cdl_fetched_at IS NOT NULL "
                  "OR m.gbooks_fetched_at IS NOT NULL)")


def _cabecera(hoja):
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.cell import WriteOnlyCell
    fondo = PatternFill("solid", fgColor="7A1F2B")
    fila = []
    for i, (t, _a) in enumerate(COLUMNAS):
        c = WriteOnlyCell(hoja, value=t)
        c.fill = fondo
        c.font = Font(bold=True, color="FFFFFF", size=10)
        c.border = Border(bottom=Side(style="thin", color="FFFFFF"))
        c.alignment = Alignment(horizontal="left" if i < 3 else "center",
                                vertical="center", wrap_text=i >= 3)
        fila.append(c)
    hoja.row_dimensions[1].height = 38
    hoja.append(fila)


def generar(todos: bool = False, ruta: str | None = None) -> dict:
    global _job
    from openpyxl import Workbook

    sello = datetime.now()
    ruta = ruta or os.path.join(CARPETA, f"metadata_libros_{sello:%Y%m%d_%H%M}.xlsx")
    _job = {"status": "running", "todos": todos, "ruta": ruta,
            "started_at": sello.isoformat(), "filas": 0, "con_peso": 0,
            "con_medidas": 0, "con_autor": 0, "sin_titulo": 0,
            "errors": [], "elapsed_s": 0}
    job = _job
    t0 = time.monotonic()
    conn = None
    try:
        conn = db.get_connection()
        cur = conn.cursor(name="metadata") if db.IS_POSTGRES else conn.cursor()
        if db.IS_POSTGRES:
            cur.itersize = 5000
        cur.execute(SQL + ("" if todos else SOLO_CON_FICHA))

        wb = Workbook(write_only=True)
        guia = wb.create_sheet("Como leerlo")
        hojas = []

        def nueva(n):
            h = wb.create_sheet("Libros" if n == 1 else f"Libros {n}")
            for i, (_t, ancho) in enumerate(COLUMNAS):
                h.column_dimensions[get_letra(i + 1)].width = ancho
            h.freeze_panes = "C2"
            _cabecera(h)
            hojas.append([h, 0])
            return hojas[-1]

        tramo = nueva(1)
        while True:
            lote = cur.fetchmany(5000)
            if not lote:
                break
            for (isbn, tit, aut, edi, col, tra, ilu, idi, enc, pag, pub,
                 peso, alto, ancho, pvp, pvpb, categ, infer, desc, img,
                 prov, gedi, hay_cdl, hay_gb, act) in lote:
                p = _numero(peso)
                al, an = _numero(alto), _numero(ancho)
                if p:
                    job["con_peso"] += 1
                if al and an:
                    job["con_medidas"] += 1
                if aut:
                    job["con_autor"] += 1
                if not tit or tit == isbn:
                    job["sin_titulo"] += 1
                origen = ("ambos" if hay_cdl and hay_gb else
                          "scrapeo" if hay_cdl else
                          "google" if hay_gb else "ninguno")
                if tramo[1] >= TOPE_HOJA:
                    tramo = nueva(len(hojas) + 1)
                tramo[0].append([
                    isbn, (tit or "")[:150], (aut or "")[:90],
                    (edi or gedi or "")[:70], (col or "")[:70],
                    (tra or "")[:70], (ilu or "")[:70], (idi or "")[:40],
                    (enc or "")[:50], _entero(pag), (pub or "")[:30],
                    p, al, an,
                    float(pvp) if pvp is not None else None,
                    float(pvpb) if pvpb is not None else None,
                    (categ or "")[:90], (infer or "")[:120],
                    int(desc or 0), (img or "")[:250],
                    (prov or "")[:120], origen,
                    act.strftime("%Y-%m-%d") if act else None,
                ])
                tramo[1] += 1
                job["filas"] += 1

        _guia(guia, job, hojas)
        wb.save(ruta)
        job["bytes"] = os.path.getsize(ruta)
        job["hojas"] = len(hojas)
        job["borrados"] = _podar(ruta)
        job["status"] = "completed"
        print(f"[Metadata] {job['filas']:,} libros, {job['con_peso']:,} con peso")
    except Exception as e:
        job["status"] = "error"
        job["errors"].append(f"{type(e).__name__}: {e}"[:300])
        print(f"[Metadata] FAIL: {e!r}")
    finally:
        if conn is not None:
            try:
                conn.close()
            except Exception:
                pass
        job["elapsed_s"] = round(time.monotonic() - t0, 2)
    return job


def _guia(hoja, job, hojas):
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.cell import WriteOnlyCell
    hoja.column_dimensions["A"].width = 32
    hoja.column_dimensions["B"].width = 94

    def f(a, b=None, neg=False, rel=None):
        fila = []
        for v in (a, b):
            c = WriteOnlyCell(hoja, value=v)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            if neg:
                c.font = Font(bold=True, size=11)
            if rel:
                c.fill = PatternFill("solid", fgColor=rel)
            fila.append(c)
        hoja.append(fila)

    n = job["filas"] or 1
    f("FICHA DE LOS LIBROS", None, True)
    f("Generado", datetime.now().strftime("%d/%m/%Y %H:%M"))
    f("")
    f("QUE HAY AQUI", None, True, "F6F3F0")
    f("Libros", f"{job['filas']:,}"
      + ("" if job["todos"] else " — solo los que tienen ficha. Con todos=true "
                                "salen los 1,3 millones del espejo."))
    f("Con peso", f"{job['con_peso']:,} ({job['con_peso']*100/n:.1f}%)")
    f("Con alto y ancho", f"{job['con_medidas']:,} ({job['con_medidas']*100/n:.1f}%)")
    f("Con autor", f"{job['con_autor']:,} ({job['con_autor']*100/n:.1f}%)")
    f("Sin titulo de verdad", f"{job['sin_titulo']:,} — llevan el ISBN por nombre")
    f("")
    f("DE DONDE SALE CADA COSA", None, True, "F6F3F0")
    f("Scrapeo de Casa del Libro", "autor, editorial, coleccion, traductor, "
                                   "ilustrador, encuadernacion, paginas, "
                                   "fecha, portada y —solo aqui— PESO, ALTO "
                                   "y ANCHO.")
    f("Google Books", "editorial, idioma y paginas, para los que no se "
                      "pudieron scrapear.")
    f("Odoo", "titulo, precio de venta, categoria y proveedores.")
    f("")
    f("LAS COLUMNAS", None, True, "F6F3F0")
    for i, (t, _a) in enumerate(COLUMNAS):
        if AYUDA.get(t):
            f(f"{get_letra(i+1)} · {t}", AYUDA[t])
    f("")
    f("UN AVISO SOBRE EL PESO", None, True, "FFF4E0")
    f("", "El peso de esta hoja NO es el de Odoo. En Odoo casi todos los "
         "libros estan a 0,00 kg; el dato bueno es el del scrapeo y es el "
         "que va aqui, ya en gramos y como numero. Para los que no lo "
         "tienen, la regla del cliente es estimar 350 g.")
