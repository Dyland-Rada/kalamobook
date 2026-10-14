"""
Los dos ficheros que hacen falta para evaluar lo del IVA.

Contexto: el 14/10/2026 Sandra detecta que en los marketplaces se cobra por
debajo del PVP del editor. La causa es que once de los quince proveedores
nos mandan la BASE IMPONIBLE en el campo donde guardamos el precio CON IVA,
y como el catalogo elige al mas barato, ese precio gana siempre.

    EJEMPLOS  -- unas pocas filas por proveedor de cada caso (con IVA y sin
                 IVA), para que cada uno pueda comprobarlo contra su propio
                 fichero sin bucear en medio millon de lineas.

    CATALOGO  -- una fila por ISBN con el PVP de CADA proveedor en su
                 columna, el PVP de referencia, lo que cobramos hoy y la
                 diferencia en euros. Es el que permite sumar el impacto.

La REFERENCIA es el feed de descuentos de AZETA, que es la unica fuente que
trae PVP y tipo de IVA en columnas separadas; de ahi se sabe cual es el
precio de verdad y cual seria sin IVA. Sin esa referencia no hay forma de
distinguir "me mandan 21,15 sin IVA" de "el libro vale 21,15".

Los dos se generan en streaming y con cursor del lado del servidor, y hay
que llamarlos desde un hilo: son minutos de trabajo sincrono.
"""
import os
import time
from datetime import datetime

import db
from informe_stock import get_letra, _nombres_proveedor, _podar

CARPETA = os.environ.get("INFORME_STOCK_DIR", "/tmp")
TOLERANCIA = 0.02

_job: dict | None = None


def get_status() -> dict:
    return dict(_job) if _job else {"status": "idle"}


def _clasifica(precio, pvp, iva):
    """'CON IVA', 'SIN IVA' o 'OTRO' comparando contra la referencia."""
    if precio is None or pvp is None or not pvp:
        return "SIN REFERENCIA"
    if abs(float(precio) - float(pvp)) <= TOLERANCIA:
        return "CON IVA"
    if iva:
        base = round(float(pvp) / (1 + float(iva) / 100.0), 2)
        if abs(float(precio) - base) <= TOLERANCIA:
            return "SIN IVA"
    return "OTRO"


def _cabecera(hoja, columnas, ancho_b=46):
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.cell import WriteOnlyCell
    fondo = PatternFill("solid", fgColor="7A1F2B")
    letra = Font(bold=True, color="FFFFFF", size=10)
    raya = Border(bottom=Side(style="thin", color="FFFFFF"))
    hoja.row_dimensions[1].height = 40
    fila = []
    for i, t in enumerate(columnas):
        c = WriteOnlyCell(hoja, value=t)
        c.fill, c.font, c.border = fondo, letra, raya
        c.alignment = Alignment(horizontal="left" if i < 2 else "center",
                                vertical="center", wrap_text=i >= 2)
        fila.append(c)
    hoja.append(fila)


def ejemplos(por_caso: int = 15, ruta: str | None = None) -> dict:
    """Unas pocas filas de cada caso y proveedor, para poder comprobarlo."""
    global _job
    from openpyxl import Workbook

    sello = datetime.now()
    ruta = ruta or os.path.join(CARPETA, f"precios_ejemplos_{sello:%Y%m%d_%H%M}.xlsx")
    _job = {"status": "running", "accion": "ejemplos", "ruta": ruta,
            "started_at": sello.isoformat(), "filas": 0, "errors": [],
            "elapsed_s": 0}
    job = _job
    t0 = time.monotonic()
    conn = None
    try:
        nombres = _nombres_proveedor()
        conn = db.get_connection()
        cur = conn.cursor(name="prec_ej") if db.IS_POSTGRES else conn.cursor()
        if db.IS_POSTGRES:
            cur.itersize = 2000
        cur.execute("""
            SELECT lp.proveedor_email, lp.isbn,
                   COALESCE(NULLIF(m.name, ''), '') AS titulo,
                   lp.precio_con_iva, f.pvp, f.iva, f.editorial
            FROM libros_proveedor lp
            JOIN azeta_feed_descuentos f
              ON f.isbn = lp.isbn AND f.pvp > 0 AND f.iva > 0
            LEFT JOIN odoo_books_mirror m ON m.barcode = lp.isbn
            WHERE lp.precio_con_iva > 0
            ORDER BY lp.proveedor_email, lp.isbn
        """)
        wb = Workbook(write_only=True)
        hoja = wb.create_sheet("Ejemplos")
        hoja.freeze_panes = "C2"
        for col, an in (("A", 26), ("B", 15), ("C", 44), ("D", 13), ("E", 13),
                        ("F", 8), ("G", 13), ("H", 12), ("I", 24)):
            hoja.column_dimensions[col].width = an
        cols = ["PROVEEDOR", "ISBN", "TITULO", "NOS MANDA", "PVP REAL",
                "IVA %", "PVP SIN IVA", "DIFERENCIA", "VEREDICTO"]
        _cabecera(hoja, cols)

        cupo: dict = {}
        for email, isbn, titulo, precio, pvp, iva, editorial in cur:
            caso = _clasifica(precio, pvp, iva)
            if caso == "SIN REFERENCIA":
                continue
            clave = (email, caso)
            if cupo.get(clave, 0) >= por_caso:
                continue
            cupo[clave] = cupo.get(clave, 0) + 1
            base = round(float(pvp) / (1 + float(iva) / 100.0), 2)
            hoja.append([
                nombres.get(email, email), isbn, titulo or editorial or "",
                float(precio), float(pvp), float(iva), base,
                round(float(precio) - float(pvp), 2), caso,
            ])
            job["filas"] += 1
        cur.close()

        hoja.auto_filter.ref = f"A1:I{job['filas'] + 1}"
        wb.save(ruta)
        job["bytes"] = os.path.getsize(ruta)
        job["borrados"] = _podar(ruta)
        job["status"] = "completed"
    except Exception as e:
        job["status"] = "error"
        job["errors"].append(f"{type(e).__name__}: {e}"[:300])
        print(f"[InfPrecios] ejemplos FAIL: {e!r}")
    finally:
        if conn is not None:
            try:
                conn.close()
            except Exception:
                pass
        job["elapsed_s"] = round(time.monotonic() - t0, 2)
    return job


def catalogo(ruta: str | None = None) -> dict:
    """Una fila por ISBN con el PVP de cada proveedor y la diferencia."""
    global _job
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.formatting.rule import CellIsRule

    sello = datetime.now()
    ruta = ruta or os.path.join(CARPETA, f"precios_catalogo_{sello:%Y%m%d_%H%M}.xlsx")
    _job = {"status": "running", "accion": "catalogo", "ruta": ruta,
            "started_at": sello.isoformat(), "isbns": 0, "filas": 0,
            "con_referencia": 0, "afectados": 0, "impacto_eur": 0.0,
            "errors": [], "elapsed_s": 0}
    job = _job
    t0 = time.monotonic()
    conn = None
    try:
        nombres = _nombres_proveedor()
        columnas = sorted(set(nombres.values()))
        indice = {n: i for i, n in enumerate(columnas)}

        conn = db.get_connection()
        cur = conn.cursor(name="prec_cat") if db.IS_POSTGRES else conn.cursor()
        if db.IS_POSTGRES:
            cur.itersize = 2000
        cur.execute("""
            SELECT lp.isbn,
                   COALESCE(NULLIF(m.name, ''), '') AS titulo,
                   f.editorial, f.pvp, f.iva,
                   c.precio_marketplace, c.proveedor,
                   lp.proveedor_email, lp.precio_con_iva, lp.stock_disponible
            FROM libros_proveedor lp
            LEFT JOIN odoo_books_mirror m ON m.barcode = lp.isbn
            LEFT JOIN azeta_feed_descuentos f ON f.isbn = lp.isbn
            LEFT JOIN catalogo_publicable c ON c.isbn = lp.isbn
            WHERE lp.precio_con_iva > 0
            ORDER BY lp.isbn, lp.proveedor_email
        """)

        wb = Workbook(write_only=True)
        hoja = wb.create_sheet("Precios")
        hoja.freeze_panes = "D2"
        fijas = ["ISBN", "TITULO", "EDITORIAL", "PVP REAL", "IVA %",
                 "PVP SIN IVA", "COBRAMOS HOY", "SIRVE", "DIFERENCIA",
                 "DIF %", "CASO"]
        for col, an in (("A", 15), ("B", 42), ("C", 24), ("D", 11), ("E", 8),
                        ("F", 12), ("G", 13), ("H", 20), ("I", 12), ("J", 9),
                        ("K", 14)):
            hoja.column_dimensions[col].width = an
        for i in range(len(columnas)):
            hoja.column_dimensions[get_letra(len(fijas) + 1 + i)].width = 13
        _cabecera(hoja, fijas + columnas)

        actual = None
        datos: dict = {}
        precios: list = []

        def volcar():
            if actual is None:
                return
            pvp = datos.get("pvp")
            iva = datos.get("iva")
            mk = datos.get("mk")
            base = (round(float(pvp) / (1 + float(iva) / 100.0), 2)
                    if pvp and iva else None)
            dif = pct = None
            caso = "SIN REFERENCIA"
            if pvp and mk:
                dif = round(float(mk) - float(pvp), 2)
                pct = round(dif / float(pvp) * 100, 2)
                if base is not None and abs(float(mk) - base) <= 0.05:
                    caso = "COBRAMOS SIN IVA"
                elif abs(dif) <= 0.05:
                    caso = "CORRECTO"
                else:
                    caso = "OTRA DIFERENCIA"
            if caso == "COBRAMOS SIN IVA":
                job["afectados"] += 1
                job["impacto_eur"] += abs(dif or 0)
            if pvp:
                job["con_referencia"] += 1
            hoja.append([
                actual, datos.get("titulo", ""), datos.get("editorial", ""),
                float(pvp) if pvp else None,
                float(iva) if iva else None, base,
                float(mk) if mk else None,
                nombres.get(datos.get("sirve"), datos.get("sirve")),
                dif, pct, caso,
            ] + precios)
            job["isbns"] += 1

        for (isbn, titulo, editorial, pvp, iva, mk, sirve,
             email, precio, _stock) in cur:
            job["filas"] += 1
            if isbn != actual:
                volcar()
                actual = isbn
                datos = {"titulo": titulo or "", "editorial": editorial or "",
                         "pvp": pvp, "iva": iva, "mk": mk, "sirve": sirve}
                precios = [None] * len(columnas)
            quien = nombres.get(email)
            if quien in indice:
                precios[indice[quien]] = float(precio)
        volcar()
        cur.close()

        ultima = get_letra(len(fijas) + len(columnas))
        hoja.auto_filter.ref = f"A1:{ultima}{job['isbns'] + 1}"
        col_caso = get_letra(len(fijas))
        hoja.conditional_formatting.add(
            f"{col_caso}2:{col_caso}{job['isbns'] + 1}",
            CellIsRule(operator="equal", formula=['"COBRAMOS SIN IVA"'],
                       fill=PatternFill("solid", fgColor="FBD5D5"),
                       font=Font(color="9B1C1C", bold=True)))
        wb.save(ruta)
        job["bytes"] = os.path.getsize(ruta)
        job["impacto_eur"] = round(job["impacto_eur"], 2)
        job["borrados"] = _podar(ruta)
        job["status"] = "completed"
        print(f"[InfPrecios] catalogo: {job['isbns']:,} ISBN, "
              f"{job['afectados']:,} afectados")
    except Exception as e:
        job["status"] = "error"
        job["errors"].append(f"{type(e).__name__}: {e}"[:300])
        print(f"[InfPrecios] catalogo FAIL: {e!r}")
    finally:
        if conn is not None:
            try:
                conn.close()
            except Exception:
                pass
        job["elapsed_s"] = round(time.monotonic() - t0, 2)
    return job
